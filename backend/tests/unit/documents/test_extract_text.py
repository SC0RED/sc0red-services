"""Tests for document text extraction."""

import io

import pytest

from src.documents.extract_text import (
    _MAX_CHARS_PER_DOCUMENT,
    combine_document_texts,
    extract_text,
)


class TestExtractText:
    def test_extract_plain_text(self):
        content = b"Hello, this is a plain text document."
        result = extract_text(content, "txt")
        assert result == "Hello, this is a plain text document."

    def test_extract_csv(self):
        content = b"name,age\nAlice,30\nBob,25"
        result = extract_text(content, "csv")
        assert "name,age" in result
        assert "Alice,30" in result

    def test_extract_markdown(self):
        content = b"# Heading\n\nSome **bold** text."
        result = extract_text(content, "md")
        assert "# Heading" in result

    def test_unsupported_type_raises(self):
        with pytest.raises(ValueError, match="Unsupported file type: pptx"):
            extract_text(b"data", "pptx")

    def test_normalises_file_type(self):
        content = b"test content"
        result = extract_text(content, ".TXT")
        assert result == "test content"

    def test_caps_at_max_chars(self):
        content = ("x" * (_MAX_CHARS_PER_DOCUMENT + 1000)).encode()
        result = extract_text(content, "txt")
        assert len(result) == _MAX_CHARS_PER_DOCUMENT

    def test_extract_pdf(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        # pypdf doesn't easily create text pages, so we test the flow with a blank PDF
        buffer = io.BytesIO()
        writer.write(buffer)
        buffer.seek(0)
        result = extract_text(buffer.read(), "pdf")
        # Blank PDF should return empty or whitespace-only text
        assert isinstance(result, str)

    def test_extract_docx(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        result = extract_text(buffer.read(), "docx")
        assert "First paragraph" in result
        assert "Second paragraph" in result

    def test_extract_xlsx(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Revenue"
        sheet["B1"] = 1000000
        sheet["A2"] = "EBITDA"
        sheet["B2"] = 250000
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        result = extract_text(buffer.read(), "xlsx")
        assert "Revenue" in result
        assert "1000000" in result

    def test_extract_xls_uses_xlsx_extractor(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "Test"
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        result = extract_text(buffer.read(), "xls")
        assert "Test" in result


class TestCombineDocumentTexts:
    def test_combines_with_separator(self):
        texts = ["Document 1 text", "Document 2 text"]
        result = combine_document_texts(texts)
        assert result == "Document 1 text\n---\nDocument 2 text"

    def test_truncates_at_limit(self):
        texts = ["x" * 20_000, "y" * 20_000]
        result = combine_document_texts(texts)
        assert len(result) <= 25_005 + len("\n[...truncated]")
        assert result.endswith("[...truncated]")

    def test_empty_list(self):
        result = combine_document_texts([])
        assert result == ""

    def test_single_document(self):
        result = combine_document_texts(["Only one doc"])
        assert result == "Only one doc"
